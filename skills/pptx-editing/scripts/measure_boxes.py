# -*- coding: utf-8 -*-
"""빌드된 pptx 에서 **PowerPoint 가 실제로 그린 텍스트 높이**를 걷어 캐시에 적는다.

왜 있나
-------
상자 높이를 «추정»하면 그 오차가 여백에 실리고, 「빌드→넘침→줄임→빌드」 루프가 시작된다.
PIL 폰트 메트릭으로 폭을 실측해도 근사는 남는다 -- CJK 금칙처리·커닝·자간을 재현할 수 없다
(2026-08-26 teacher 포스터 잔차 1.4cm).

PowerPoint 는 정답을 안다: `TextFrame2.TextRange.BoundHeight`. 이 스크립트는 COM 을 한 번
열어 그걸 걷는다. 도형 이름에 `poster_kit.box_key()` 해시가 새겨져 있으므로 텍스트를 다시
맞춰볼 필요가 없다 -- 이름이 곧 키다.

쓰는 법 (2회 수렴, 결정론적)
---------------------------
    python 03_build.py                    # 1회차: 추정으로 그린다
    python measure_boxes.py <out.pptx>    # COM 1회로 실측을 캐시에 적는다
    python 03_build.py                    # 2회차: 캐시를 써서 «정확»하다
    python measure_boxes.py <out.pptx>    # 검산: 오차가 0 으로 수렴했는지 본다

빌드가 찍는 `PK.cache_report()` 가 100% 실측이라고 말할 때까지 돌린다. 캐시는 프로젝트
밖(`agent/cache/pptx_box_heights.json`)에 살아서 **다음 포스터·덱은 1회차부터 정확**하다.

읽기 전용이다
-------------
`ReadOnly=True` 로 열고 저장하지 않는다 -- 렌더가 pptx 를 오염시킨 사고가 2026-08-26 하루에
두 번 났다(도형이 두 벌씩 복제). 그래도 뒤에 `build_guard.py verify` 로 md5 를 확인하는 편이 낫다.
"""

import argparse
import io
import json
import os
import sys

PT2CM = 2.54 / 72.0


def _pk():
    """옆의 poster_kit 에서 상수를 «가져온다» -- 복사본을 두면 언젠가 갈리고,
    그때 캐시 전체가 그 차이만큼 조용히 틀린다."""
    import importlib.util
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), "poster_kit.py")
    spec = importlib.util.spec_from_file_location("poster_kit", p)
    m = importlib.util.module_from_spec(spec)
    sys.modules["poster_kit"] = m
    spec.loader.exec_module(m)
    return m


_PK = _pk()
BOX_PAD_CM = _PK.BOX_PAD_CM        # 캐시가 담는 값 = 텍스트 높이 + 이 패드
CACHE_DEFAULT = _PK.CACHE_PATH


def collect(path):
    """(키 -> 실측높이cm, 키 -> 빌드가 예측한 상자높이cm) 를 돌려준다."""
    import win32com.client as win32
    from pptx import Presentation

    # 파일에 «적힌» 표 높이 -- 빌더가 의도한 값이다. COM 으로 열면 이미 갱신돼 있으므로
    # 반드시 열기 «전에» 여기서 읽어 둔다.
    # ⚠ 키는 «(슬라이드, shape_id)» 다. shape_id 는 **슬라이드 안에서만 고유**해서, 슬라이드
    #   번호를 빼면 덱에서 서로 다른 표가 같은 칸을 덮어쓴다 -- 1슬라이드 포스터에서는 절대
    #   안 드러나고 덱에서만 거짓 양성으로 나온다(2026-08-27, 실제로 그렇게 오보고했다).
    stored, stored_rows = {}, {}
    for _si, _sl in enumerate(Presentation(os.path.abspath(path)).slides, 1):
        for _sh in _sl.shapes:
            if _sh.has_table:
                stored[(_si, _sh.shape_id)] = _sh.height / 360000.0
                stored_rows[(_si, _sh.shape_id)] = [round(r.height / 360000.0, 4)
                                                    for r in _sh.table.rows]

    app = win32.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(os.path.abspath(path), WithWindow=False, ReadOnly=True)
    measured, predicted, grown = {}, {}, []
    try:
        for si in range(1, pres.Slides.Count + 1):
            sl = pres.Slides(si)
            for sp in sl.Shapes:
                # ── 표: 행 «자동확장»을 잡는다 ────────────────────────────────
                # PowerPoint 는 셀이 넘치면 그 행만 늘린다. 표가 커졌다는 것은 곧
                # **빌더가 받은 높이가 거짓이고 그 아래 놓은 것이 겹쳤다**는 뜻이다.
                #
                # ⚠ 비교 대상을 틀리기 쉽다. PowerPoint 는 행을 늘리면서 «행 높이도 같이»
                #   갱신하므로 열어 본 뒤에는 행합 == 전체높이라 언제나 일치한다. 참인 비교는
                #   **파일에 적힌 값(python-pptx) vs 렌더된 값(COM)** 이다.
                try:
                    if int(sp.HasTable) == -1:
                        want = stored.get((si, int(sp.Id)))
                        real = float(sp.Height) * PT2CM
                        if want and real - want > 0.05:
                            grown.append((si, str(sp.Name), want, real))
                        # 킷이 만든 표면 **행별** 실측을 캐시에 담는다. 표는 총높이만으로는
                        # 행을 배치할 수 없으므로 값이 스칼라가 아니라 리스트다.
                        # ⚠ `Rows(r).Height` 는 잘 읽힌다 -- 앞서 「0.000 이 나온다」고 본 것은
                        #   반복문 안에서 app.Quit() 을 불러 «죽은» COM 인스턴스를 잡았기
                        #   때문이었다(2026-08-27 확인). 늘어난 행도 늘어난 값으로 나온다.
                        nm = str(sp.Name)
                        if nm.startswith("pk:"):
                            tb = sp.Table
                            measured[nm] = [round(float(tb.Rows(r).Height) * PT2CM, 4)
                                            for r in range(1, tb.Rows.Count + 1)]
                            predicted[nm] = stored_rows.get((si, int(sp.Id)), [])
                        continue
                except Exception:
                    pass
                name = str(sp.Name)
                if not name.startswith("pk:"):
                    continue          # 킷이 안 만든 도형 -- 키가 없으니 캐시할 수 없다
                try:
                    if not (sp.HasTextFrame and sp.TextFrame.HasText):
                        continue
                    bh = float(sp.TextFrame2.TextRange.BoundHeight) * PT2CM
                except Exception:
                    continue
                measured[name] = round(bh + BOX_PAD_CM, 4)
                predicted[name] = round(float(sp.Height) * PT2CM, 4)
    finally:
        pres.Close()
        app.Quit()
    return measured, predicted, grown


def main():
    ap = argparse.ArgumentParser(description="pptx 상자 높이 실측 -> 캐시")
    ap.add_argument("pptx")
    ap.add_argument("--cache", default=CACHE_DEFAULT)
    ap.add_argument("--dry-run", action="store_true", help="캐시에 쓰지 않고 오차만 본다")
    a = ap.parse_args()

    measured, predicted, grown = collect(a.pptx)

    # 표 경고는 캐시와 무관하게 «항상» 먼저 낸다 -- 이건 레이아웃이 이미 깨졌다는 신호다.
    for si, nm, want, real in grown:
        print("표가 늘어났다: 슬라이드 %d, %s  지정 %.2fcm -> 실제 %.2fcm (%+.2f)"
              % (si, nm, want, real, real - want))
        print("  => 빌더가 받은 표 높이가 거짓이다. 그 아래 놓은 캡션·다음 블록이 겹쳐 있다.")
        print("     셀 하나가 열 폭에서 줄바꿈된 것이니, 행 높이를 재서 정하라"
              " (pptx_kit.dtable / poster_kit.ptable 은 그렇게 한다).")
    if not measured:
        if grown:
            return 1                      # 표 경고는 위에서 이미 냈다
        print("키가 새겨진 텍스트 상자가 없다 -- 킷의 bullets()/caption() 으로 만든 상자가"
              " 있어야 캐시할 수 있다. (표는 따로 검사했고, 늘어난 표는 없다.)")
        return 0

    try:
        with io.open(a.cache, encoding="utf-8") as f:
            cache = json.load(f)
    except Exception:
        cache = {}

    def _tot(v):
        """상자는 스칼라(cm), 표는 행높이 리스트 -- 오차는 «총높이»로 비교한다."""
        return sum(v) if isinstance(v, list) else (v or 0.0)

    def _err(k):
        return _tot(predicted.get(k)) - _tot(measured[k])

    new = sum(1 for k in measured if k not in cache)
    worst_k = max(measured, key=lambda k: abs(_err(k)))
    worst = abs(_err(worst_k))
    total_err = sum(_err(k) for k in measured)
    n_tab = sum(1 for v in measured.values() if isinstance(v, list))

    print("도형 %d개 실측(표 %d개 포함) · 캐시 신규 %d개" % (len(measured), n_tab, new))
    print("예측 오차: 최대 %.2fcm · 합계 %+.2fcm (이 합이 열 끝 y 의 오차다)" % (worst, total_err))
    if worst > 0.05:
        print("  최악: %s  예측 %.2f vs 실측 %.2f"
              % (worst_k, _tot(predicted.get(worst_k)), _tot(measured[worst_k])))
    if worst <= 0.05:
        print("  수렴 -- 다시 빌드해도 같은 레이아웃이 나온다.")

    if a.dry_run:
        print("(dry-run: 캐시를 쓰지 않았다)")
        return 0

    cache.update(measured)
    d = os.path.dirname(os.path.abspath(a.cache))
    if not os.path.isdir(d):
        os.makedirs(d)
    with io.open(a.cache, "w", encoding="utf-8", newline="\n") as f:
        f.write(json.dumps(cache, ensure_ascii=False, indent=0, sort_keys=True))
    print("캐시 %d개 항목 -> %s" % (len(cache), a.cache))
    print("다시 빌드하면 이 상자들은 실측 높이로 그려진다.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
