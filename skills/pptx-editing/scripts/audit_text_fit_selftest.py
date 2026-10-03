# -*- coding: utf-8 -*-
"""audit_text_fit_selftest.py — 「상자 밖」 검사가 «걸리는 입력»과 «안 걸리는 입력» 둘 다에서 맞나.

    python audit_text_fit_selftest.py

2026-09-28 에 이 검사를 넣은 계기: 카드(테두리 있는 사각형) «안»에 놓인 글상자가 카드를 뚫고
나갔는데 감사가 통과시켰다. 이웃이 글을 안 가진 도형이면 «내용을 지닌 이웃»이 아니라고 보고
건너뛰었기 때문이다. 배경 위에 얹힌 평범한 넘침과 가르는 기준은 «그 도형이 글상자를 담고 있는가»다.

⚠ 이 시험은 PowerPoint COM 을 쓴다(감사 자체가 그렇다). 없으면 건너뛴다고 말하고 0 을 낸다.
"""
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

LONG = ("대마 사용은 성별·나이·인종·교육·산업·직업에 따라 크게 다르고, 손상·사망 위험이 큰 "
        "일부 업종에서 평균보다 높았으며, 쓰는 사람 가운데 매일 쓰는 비율도 높은 편이다.")


def fixture(path):
    """1장 = 카드를 뚫고 나가는 글, 2장 = 같은 글이 충분한 카드 안에 들어간 경우."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for h in (1.2, 3.2):                      # 작은 카드(걸려야 한다) · 넉넉한 카드(걸리면 안 된다)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5),
                                 Inches(6.0), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        box.line.color.rgb = RGBColor(0xC8, 0xC8, 0xC8)
        tb = s.shapes.add_textbox(Inches(0.95), Inches(1.65), Inches(5.7), Inches(h - 0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = LONG
        tf.paragraphs[0].font.size = Pt(20)

    # 3·4장: 그림 위 글상자. 작은 그림 위는 겹침이고, 슬라이드를 덮는 «배경» 그림 위는
    # 정상 레이어링이다(원본 쪽 이미지를 깔고 여백에 주석을 얹는 덱 — islr2-45 보고).
    import struct
    import zlib

    def _png(w, h):
        def chunk(tag, data):
            c = tag + data
            return struct.pack(">I", len(data)) + c + struct.pack(">I", zlib.crc32(c))
        raw = b"".join(b"\x00" + b"\xdd" * (w * 3) for _ in range(h))
        return (b"\x89PNG\r\n\x1a\n"
                + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))

    import tempfile as _tf
    png = Path(_tf.gettempdir()) / "_fit_probe.png"
    png.write_bytes(_png(40, 30))
    for pic_w, pic_h in ((4.0, 3.0), (13.0, 7.2)):        # 작은 그림 · 배경 그림
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(png), Inches(0.2), Inches(0.2), Inches(pic_w), Inches(pic_h))
        tb = s.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(2.0), Inches(0.6))
        tb.text_frame.paragraphs[0].text = "그림 위에 얹은 주석"
        tb.text_frame.paragraphs[0].font.size = Pt(14)
    prs.save(str(path))


def run():
    try:
        import win32com.client                                    # noqa: F401
        import pptx                                               # noqa: F401
    except Exception:                                             # noqa: BLE001
        print("PowerPoint COM(pywin32) 또는 python-pptx 가 없다 — 건너뛴다")
        return 0
    with tempfile.TemporaryDirectory() as td:
        deck = Path(td) / "fit_fixture.pptx"
        fixture(deck)
        r = subprocess.run([sys.executable, str(HERE / "audit_text_fit.py"), str(deck)],
                           capture_output=True, text=True, encoding="utf-8", errors="replace")
        out = (r.stdout or "") + (r.stderr or "")
    bad = []
    if "[상자 밖" not in out:
        bad.append("작은 카드를 뚫고 나간 글을 «못 잡았다» — 검사가 죽었다")
    if "슬  2" in out.split("[상자 밖", 1)[-1].split("\n\n", 1)[0]:
        bad.append("넉넉한 카드(2장)를 «잘못 잡았다» — 오탐")

    seg = out.split("[자리 겹침", 1)
    if len(seg) == 1:
        bad.append("작은 그림 위의 글상자를 «못 잡았다» — 자리 겹침 검사가 죽었다")
    else:
        block = seg[1].split(chr(10) + chr(10), 1)[0]
        if "슬  3" not in block:
            bad.append("작은 그림 위의 글상자를 «못 잡았다»")
        if "슬  4" in block:
            bad.append("배경 그림 위의 글상자를 «잘못 잡았다» — islr2-45 가 보고한 오탐")
    print("4 cases (2 of them must NOT fire)")
    if bad:
        print("FAIL %d:" % len(bad))
        for b in bad:
            print("  " + b)
        print("--- 감사 출력 ---")
        print(out[:1200])
        return 1
    print("  OK  all pass")
    return 0


if __name__ == "__main__":
    sys.exit(run())
