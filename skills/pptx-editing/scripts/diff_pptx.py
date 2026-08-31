# -*- coding: utf-8 -*-
"""diff_pptx.py — what changed between two decks. Run this BEFORE you overwrite a deck
someone hand-edited, and again AFTER you fold their edits into the build script.

    python diff_pptx.py OLD.pptx NEW.pptx
    python diff_pptx.py shipped.zip::decks/week2.pptx current.pptx   # baseline inside a zip
    python diff_pptx.py OLD.pptx NEW.pptx --text-only

SKILL.md tells you to look before you clobber a hand-edited deck. This is the looking.
It is also the check that your fold-in worked: rebuild, diff against their file, and the
only differences left should be the ones you meant to introduce.

⚠ THE TRAP THIS EXISTS FOR — compare PARAGRAPHS, never RUNS.
PowerPoint re-splits every paragraph into word-level runs when it saves. A paragraph
nobody touched comes back as 40 new runs, so a run-level diff reports the whole deck as
changed and buries the four real edits. (Measured: a 39-slide deck with 5 hand-edited
slides produced ~600 spurious run rows; at paragraph level it produced 11 true rows.)
Runs also lose their explicit size/colour on re-save, so per-run formatting comparison
invents changes. This script joins runs per paragraph and takes formatting from the first.

What it reports, per slide:
  · text of a paragraph or table row changed / added / removed
  · first-run formatting: size, bold, colour        (catches "grey caption -> red bold")
  · shape geometry: left/top/width/height in inches (catches "table tightened", "box widened")
  · shapes added or removed, and speaker-note changes

What it deliberately does NOT report: run splitting, XML attribute churn, zip timestamps.
Those differ on every PowerPoint save and mean nothing.

Self-test:  python diff_pptx.py --selftest
"""
import argparse
import io
import os
import sys
import zipfile

if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from pptx import Presentation

EMU = 914400.0


def _open(spec):
    """Path, or ARCHIVE.zip::member-inside-it (the shipped-baseline case)."""
    if "::" in spec:
        arc, member = spec.split("::", 1)
        z = zipfile.ZipFile(arc)
        names = [n for n in z.namelist() if n == member or n.endswith("/" + member)]
        if not names:
            # zip entries written on a CP949 box are not UTF-8 flagged; match loosely
            names = [n for n in z.namelist()
                     if member.split("/")[-1] in n or n.endswith(".pptx") and member in n]
        if not names:
            raise SystemExit("zip 안에서 못 찾음: %s (안에 있는 pptx: %s)"
                             % (member, [n for n in z.namelist() if n.endswith(".pptx")][:5]))
        return Presentation(io.BytesIO(z.read(names[0])))
    return Presentation(spec)


def _fmt(run):
    col = None
    try:
        col = str(run.font.color.rgb)
    except Exception:
        pass
    return (run.font.size.pt if run.font.size else None, run.font.bold, col)


def snapshot(prs, text_only=False):
    """{key: (shape name, text, size, bold, colour, geometry)} — paragraph granularity."""
    out = {}
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            geo = None
            if not text_only:
                try:
                    geo = tuple(round(v / EMU, 2) for v in
                                (sh.left, sh.top, sh.width, sh.height))
                except TypeError:          # a placeholder can carry None
                    geo = None
                # 표·그림인지 표시해 둔다 — 텍스트 상자의 «높이만» 달라진 것은 대개
                # PowerPoint 의 shrink-wrap 이지 사람이 고친 것이 아니다(아래 report 참조).
                kind = "table" if sh.has_table else ("text" if sh.has_text_frame else "other")
                out[(i, sh.shape_id, "geo")] = (sh.name, kind, None, None, None, geo)
            if sh.has_text_frame:
                for pi, p in enumerate(sh.text_frame.paragraphs):
                    txt = "".join(r.text for r in p.runs)
                    if not txt.strip():
                        continue
                    f = _fmt(p.runs[0]) if not text_only else (None, None, None)
                    out[(i, sh.shape_id, "p%d" % pi)] = (sh.name, txt) + f + (None,)
            if sh.has_table:
                for ri, row in enumerate(sh.table.rows):
                    cells = " || ".join(c.text.strip() for c in row.cells)
                    out[(i, sh.shape_id, "r%d" % ri)] = (
                        sh.name, cells, None, None, None, None)
        if s.has_notes_slide:
            t = (s.notes_slide.notes_text_frame.text or "").strip()
            if t:
                out[(i, -1, "note")] = ("(speaker note)", t, None, None, None, None)
    return out


LABEL = {2: "글자", 3: "크기", 4: "굵게", 5: "색", 6: "좌표"}


def report(a, b):
    slides = sorted({k[0] for k in set(a) | set(b)})
    total = 0
    for sl in slides:
        ka = {k: v for k, v in a.items() if k[0] == sl}
        kb = {k: v for k, v in b.items() if k[0] == sl}
        lines = []
        for k in sorted(set(ka) | set(kb), key=str):
            x, y = ka.get(k), kb.get(k)
            if x == y:
                continue
            if x is None:
                lines.append("  + 새로 생김  %-16s %r" % (y[0][:16], y[1][:70]))
            elif y is None:
                lines.append("  - 사라짐    %-16s %r" % (x[0][:16], x[1][:70]))
            else:
                for idx in (1, 2, 3, 4, 5):
                    if x[idx] != y[idx]:
                        lab = LABEL[idx + 1]
                        if idx == 1:
                            lines.append("  ~ %-16s %s\n      %r\n   -> %r"
                                         % (x[0][:16], lab, x[1][:70], y[1][:70]))
                        else:
                            hint = ""
                            # 텍스트 상자에서 «높이만» 달라졌으면 PowerPoint 가 저장하며
                            # 내용에 맞춰 줄인 것일 가능성이 높다. 글자는 그대로이고 상자는
                            # 자동으로 늘고 주므로 렌더 결과가 같다 — 지우지는 말고 «표시»만.
                            if (idx == 5 and x[1] == "text" and x[5] and y[5]
                                    and x[5][:3] == y[5][:3]):
                                hint = "   ← 높이만: PowerPoint 자동 맞춤일 수 있다"
                            lines.append("  ~ %-16s %-4s %r -> %r   (%s)%s"
                                         % (x[0][:16], lab, x[idx], y[idx], x[1][:26], hint))
        if lines:
            total += len(lines)
            print("=" * 12, "슬라이드", sl)
            for ln in lines:
                print(ln)
    print("\n달라진 항목 %d건" % total)
    return total


def _selftest():
    """Prove the paragraph-vs-run point and the three change kinds, without a template."""
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    def make(text, size, bold, colour, width, split_runs):
        p = Presentation()
        sl = p.slides.add_slide(p.slide_layouts[6])
        box = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(width), Inches(1))
        para = box.text_frame.paragraphs[0]
        chunks = text.split(" ") if split_runs else [text]
        for j, c in enumerate(chunks):
            r = para.add_run()
            r.text = c if j == len(chunks) - 1 else c + " "
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(colour)
        return p

    base = make("한 문장 그대로 둔다", 18, False, "1A1A1A", 5, split_runs=False)
    same_but_split = make("한 문장 그대로 둔다", 18, False, "1A1A1A", 5, split_runs=True)
    changed = make("한 문장 그대로 둔다", 18, True, "C00000", 6, split_runs=True)

    ok = []
    n = report(snapshot(base), snapshot(same_but_split))
    ok.append(("run 쪼개기만 다르면 «차이 없음»", n == 0))
    n = report(snapshot(base), snapshot(changed))
    ok.append(("굵기·색·폭 변경은 잡는다", n == 3))
    n = report(snapshot(base, text_only=True), snapshot(changed, text_only=True))
    ok.append(("--text-only 는 서식·좌표를 무시한다", n == 0))

    print()
    bad = 0
    for name, good in ok:
        print("  %s  %s" % ("PASS" if good else "FAIL", name))
        bad += 0 if good else 1
    print("\n%d PASS / %d FAIL" % (len(ok) - bad, bad))
    return 1 if bad else 0


def main():
    if "--selftest" in sys.argv:
        return _selftest()
    ap = argparse.ArgumentParser()
    ap.add_argument("old", help="baseline .pptx, or ARCHIVE.zip::member")
    ap.add_argument("new", help="current .pptx")
    ap.add_argument("--text-only", action="store_true",
                    help="ignore formatting and geometry; compare wording only")
    a = ap.parse_args()
    for spec in (a.old, a.new):
        if "::" not in spec and not os.path.exists(spec):
            raise SystemExit("파일이 없습니다: %s" % spec)
    n = report(snapshot(_open(a.old), a.text_only),
               snapshot(_open(a.new), a.text_only))
    return 1 if n else 0


if __name__ == "__main__":
    sys.exit(main())
