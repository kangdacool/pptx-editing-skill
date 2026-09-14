# -*- coding: utf-8 -*-
"""deck_mono — 발표·절차 덱의 «흑백» 스타일. 한 곳에만 둔다.

    # 어느 폴더에서 돌리든 «스킬 경로»로 찾는다 — 스크립트에 절대경로를 박지 않는다.
    import os, sys
    sys.path.insert(0, os.path.join(os.path.expanduser("~"), ".claude",
                                    "skills", "pptx-editing", "scripts"))
    from deck_mono import *
    prs = new_deck(); set_running("문서 이름")
    s = cover(prs, "머리말", "제목", "부제")
    s = content(prs, "제목", "절 이름");  bullets(s, [...]);  card(s, ...)
    rich(p, "독학 18명(90.0%)", sz=18)   # 수치 토막만 굵게 — 흑백 덱의 유일한 강조 수단

왜 여기 있나
    같은 흑백 스타일이 세 번째 덱에서 또 필요해졌다(실습 덱 · 저널클럽 덱 · 절차서).
    규칙: **세 번째면 복사하지 말고 공용으로 올린다** — 복사한 헬퍼는 반드시 갈라진다.
    기계 규격(템플릿 로드·이미지 맞춤·발표자 노트)은 같은 폴더의 pptx_kit 이 정본이고,
    여기엔 «색과 문법»만 둔다.

    ⭐ 왜 개인 도구 폴더가 아니라 «스킬»인가
    덱과 빌드 스크립트를 주제별 폴더로 나누자 상대경로 import 가 깨졌다. 스킬 경로는
    홈 디렉터리에서 풀리므로 «어느 폴더에 두든» 찾을 수 있다 — 그게 여기여야 하는 이유다.

    ⚠ 이 스타일의 먼저 있던 구현들은 그대로 둔다 — 살아 있는 산출물을 만들고 있어서 지금
    갈아끼우면 그 덱들이 흔들린다. 다음에 그 덱을 손볼 때 이쪽으로 옮긴다.

팔레트 근거
    실제로 배포한 실습 덱에서 실측한 값. 색은 «자료»의 것이고 덱은 흑백이다 —
    발표덱에는 논문 원본 그림이 원색으로 들어오므로, 덱에 색이 있으면 충돌하고
    청중이 색을 «의미»로 읽는다. 강조는 ACCENT 하나뿐이고 덱당 한둘만 쓴다.

한글 크기 하한(ppt_rules): 제목 26 · 본문 16 · 부연 15 · **13 미만 절대 금지.**
"""
import os
import re as _re
import sys

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# ── 기계 규격은 스킬에서 가져온다(복사하지 않는다) ──
def _skill_scripts_dir():
    env = os.environ.get("PPTX_KIT")
    if env:
        return os.path.dirname(env) if env.lower().endswith(".py") else env
    base = os.environ.get("CLAUDE_SKILLS_DIR") or os.path.join(
        os.path.expanduser("~"), ".claude", "skills")
    return os.path.join(base, "pptx-editing", "scripts")


_scripts = _skill_scripts_dir()
if not os.path.isfile(os.path.join(_scripts, "pptx_kit.py")):
    raise ImportError(f"pptx_kit.py 를 {_scripts} 에서 못 찾았다 — pptx-editing 스킬을 설치하거나 "
                      "PPTX_KIT 로 경로를 지정할 것")
sys.path.insert(0, _scripts)
from pptx_kit import (wrapped_lines,  # noqa: E402
                      new_deck as _new_deck, blank_slide_layout as _blank,
                      rect, slide_number as _pageno, speaker_note, fit_picture)

# ── 팔레트 ──
INK      = RGBColor(0x00, 0x00, 0x00)   # 제목
BODY     = RGBColor(0x1A, 0x1A, 0x1A)   # 본문
MUTE     = RGBColor(0x59, 0x59, 0x59)   # 부연·캡션
HDR_L    = RGBColor(0xA6, 0xA6, 0xA6)   # 러닝헤더 좌
HDR_R    = RGBColor(0x7F, 0x7F, 0x7F)   # 러닝헤더 우
PAGENO   = RGBColor(0x89, 0x89, 0x89)
HAIRLINE = RGBColor(0xC8, 0xC8, 0xC8)
BAND     = RGBColor(0xED, 0xED, 0xED)   # 카드 머리띠
CODEBG   = RGBColor(0xF7, 0xF7, 0xF7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT   = RGBColor(0xC0, 0x00, 0x00)   # 아껴 쓴다 — 덱당 한둘

KO, MONO = "맑은 고딕", "Consolas"
SW_IN, SH_IN = 13.3333, 7.5
_RUNNING = ""


def set_running(text):
    """모든 장 좌상단에 깔 러닝헤더(문서·주차 이름)."""
    global _RUNNING
    _RUNNING = text


def set_font(name):
    """본문 글꼴을 바꾼다(영문 덱이면 "Arial"). 기본은 맑은 고딕.

    ⚠ `deck_mono.KO = "Arial"` 로 «모듈 전역만» 바꾸면 안 된다 — tb() 의 기본 인자는 def 시점에
    묶여 맑은 고딕으로 남고, content()·table() 이 만드는 글자가 전부 그쪽을 탄다. bullets·card·
    rich 는 호출 시점에 읽으므로 바뀌어서, **«일부만» 바뀌는 더 나쁜 상태**가 된다. 그래서 tb 는
    font=None 을 받아 호출 시점에 KO 를 읽고, 바꾸는 길은 이 함수 하나로 둔다.
    """
    global KO
    KO = name


def _flat(sh):
    """테마 기본 그림자를 끈다 — 흑백 덱에서 상자·괘선이 흐릿하게 번진다.

    이 스타일의 첫 구현에는 있던 규칙인데 공용으로 올릴 때 빠졌다 — 「색과 문법」만 옮기면
    이런 한 줄이 소리 없이 사라진다. 카드·괘선·표가 전부 번진 뒤에야 드러난다.
    """
    try:
        sh.shadow.inherit = False
    except (AttributeError, NotImplementedError):
        pass
    return sh


_rect_raw = rect


def rect(s, l, t, w, h, color):
    """pptx_kit.rect + 그림자 끄기. 이 모듈의 괘선·머리띠는 전부 이쪽을 탄다."""
    return _flat(_rect_raw(s, l, t, w, h, color))


def tb(s, l, t, w, h, text, sz=16, bold=False, color=BODY, align=PP_ALIGN.LEFT, font=None):
    font = font or KO                   # ⚠ 기본 인자로 두면 set_font() 가 안 먹는다
    box = _flat(s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size, p.font.bold, p.font.name = Pt(sz), bold, font
    p.font.color.rgb, p.alignment = color, align
    return box


# 본문에서 굵게 올릴 토막. 긴 것부터 맞춘다(「18명(90.0%)」이 「18명」으로 쪼개지지 않게).
# ⛔ 소수점 없는 「5점」은 넣지 않는다 — 「5점 이하」의 5점은 «기준값»이지 결과가 아니다
#    (2026-09-11 실측: 그것까지 굵어져 위계가 뒤집혔다).
EMPH = _re.compile(
    r"\d+명\(\d+(?:\.\d+)?%\)"      # 18명(90.0%)
    r"|\d+\.\d+점"                    # 4.50점
    r"|N\s*=\s*\d+"                   # N = 20
    r"|\d+명"                          # 20명
    r"|\d+(?:\.\d+)?%"                # 90.0%
)


def rich(p, text, *, sz, color=BODY, bold=False):
    """수치 토막만 굵게(+검정) 올려 한 문단을 여러 run 으로 쓴다.

    ⚠ 흑백 덱의 강조는 «굵기»다 — 색은 ACCENT 하나뿐이고 그것은 급소 카드 전용이다.
    ⚠ 줄 «전체»를 굵히지 않는다. 그러면 강조가 아니라 소음이 된다.
    ⛔ 인용문과 회색 부연에는 걸지 않는다 — 인용은 남의 말이고, 부연은 일부러 내려둔 줄이라
       굵히면 위계가 뒤집힌다.
    """
    pos = 0
    for m in EMPH.finditer(text):
        if m.start() > pos:
            r = p.add_run(); r.text = text[pos:m.start()]
            r.font.size, r.font.color.rgb, r.font.bold, r.font.name = Pt(sz), color, bold, KO
        r = p.add_run(); r.text = m.group()
        r.font.size, r.font.color.rgb, r.font.bold, r.font.name = Pt(sz), INK, True, KO
        pos = m.end()
    if pos < len(text) or pos == 0:
        r = p.add_run(); r.text = text[pos:]
        r.font.size, r.font.color.rgb, r.font.bold, r.font.name = Pt(sz), color, bold, KO


def _bg(s, color=WHITE):
    r = _flat(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW_IN), Inches(SH_IN)))
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()


def running_header(s, right=""):
    if _RUNNING:
        tb(s, 0.72, 0.13, 8.4, 0.36, _RUNNING, sz=16, color=HDR_L)
    if right:
        tb(s, 8.0, 0.13, 4.6, 0.36, right, sz=16, color=HDR_R, align=PP_ALIGN.RIGHT)


def cover(prs, kicker, title, subtitle=""):
    s = prs.slides.add_slide(_blank(prs))
    _bg(s)
    tb(s, 0.72, 0.13, 11.9, 0.36, kicker, sz=16, color=HDR_L)
    tb(s, 0.9, 2.55, 11.6, 1.5, title, sz=44, color=INK, align=PP_ALIGN.CENTER)
    if subtitle:
        tb(s, 0.9, 4.35, 11.6, 1.0, subtitle, sz=20, color=BODY, align=PP_ALIGN.CENTER)
    return s


def section(prs, roman, title, subtitle=""):
    """⛔ 로마숫자를 크게 «따로» 놓지 말 것 — 흑백에서 「I」이 세로줄 하나로 보인다."""
    s = prs.slides.add_slide(_blank(prs))
    _bg(s)
    running_header(s)
    rect(s, 6.07, 2.95, 1.2, 0.035, INK)
    tb(s, 0.9, 3.25, 11.6, 0.9, f"{roman}. {title}", sz=40, bold=True, color=INK,
       align=PP_ALIGN.CENTER)
    if subtitle:
        tb(s, 0.9, 4.20, 11.6, 0.7, subtitle, sz=19, color=MUTE, align=PP_ALIGN.CENTER)
    return s


def content(prs, title, breadcrumb=""):
    s = prs.slides.add_slide(_blank(prs))
    _bg(s)
    running_header(s, breadcrumb)
    tb(s, 0.70, 0.62, 12.0, 0.66, "▪ " + title, sz=26, bold=True, color=INK)
    rect(s, 0.72, 1.40, 11.9, 0.02, HAIRLINE)
    # ⚠ 13 미만으로 내리지 말 것 — 이 파일 머리말의 하한이고 audit_font_sizes 의
    #   slide-ko 캡션 바닥값이다. 질의응답에서 청중이 부르는 것이 쪽번호다.
    _flat(_pageno(prs, s, color=PAGENO, font=KO, size=13) or s.shapes[-1])
    return s


def bullets(s, items, l=0.78, t=1.72, w=11.8, h=5.2, sz=18, gap=9):
    """items: 문자열, 또는 (True, "소제목") 머리행.  반환 = 마지막 줄 «아래 y».

    ⭐ 반환값이 card/table/picture 와 같은 바닥 y 다 — 아래에 표·강조줄을 놓을 때 이 값을 쓴다.
    2026-09-12 까지는 도형을 돌려줬고, 그래서 불릿 아래 좌표를 매번 손으로 짐작하다가 «두 번»
    글자가 겹쳤다. `audit_text_fit` 은 배경 없는 글자끼리의 겹침을 세지 않아 렌더에서만 보인다.
    (도형 자체가 필요하면 호출 직후 `s.shapes[-1]`.)

    높이는 넣은 항목 «수»가 아니라 되감긴 줄 수로 잰다 — 긴 항목은 두세 줄로 접힌다.
    """
    box = _flat(s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    line_h = sz * 0.0178 + 0.055          # deck_mono.table 과 같은 한 줄 높이
    bottom = t
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        if isinstance(it, tuple) and it and it[0] is True:
            r = p.add_run()
            r.text = it[1]
            r.font.size, r.font.bold, r.font.name = Pt(sz), True, KO
            r.font.color.rgb = INK
            if not first:
                p.space_before = Pt(gap + 5)
                bottom += (gap + 5) / 72.0
            bottom += wrapped_lines(it[1], w, sz, KO, bold=True) * line_h
        else:
            r1 = p.add_run()
            r1.text = "•  "
            r1.font.size, r1.font.name, r1.font.color.rgb = Pt(sz), KO, INK
            r2 = p.add_run()
            r2.text = it
            r2.font.size, r2.font.name, r2.font.color.rgb = Pt(sz), KO, BODY
            pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
            width = int(Pt(sz) * 0.95)
            pPr.set("marL", str(width))
            pPr.set("indent", str(-width))
            if not first:
                p.space_before = Pt(gap)
                bottom += gap / 72.0
            # 글머리표·내어쓰기만큼 실제 글자 폭이 줄어든다
            bottom += wrapped_lines(it, w - Pt(sz).inches * 0.95, sz, KO) * line_h
        first = False
    return bottom


def card(s, l, t, w, h, header, lines, sz=16, accent=False):
    """흰 상자 + 얇은 테두리 + 회색 머리띠. accent=True 면 왼쪽 빨간 세로띠(덱당 한둘)."""
    box = _flat(s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   Inches(l), Inches(t), Inches(w), Inches(h)))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = HAIRLINE
    box.line.width = Pt(1.0)
    hb = rect(s, l, t, w, 0.42, BAND)
    tfh = hb.text_frame
    tfh.vertical_anchor = MSO_ANCHOR.MIDDLE
    tfh.margin_left = Inches(0.16)
    ph = tfh.paragraphs[0]
    ph.text = header
    ph.font.size, ph.font.bold, ph.font.name = Pt(sz + 1), True, KO
    ph.font.color.rgb = INK
    if accent:
        rect(s, l, t, 0.06, h, ACCENT)
    bx = _flat(s.shapes.add_textbox(Inches(l + 0.16), Inches(t + 0.52),
                                    Inches(w - 0.32), Inches(h - 0.64)))
    tf = bx.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.font.size, p.font.name, p.font.color.rgb = Pt(sz), KO, BODY
        if i:
            p.space_before = Pt(4)
    return t + h                                    # ⬅ 아래에 무엇을 놓을 때 «이 값»을 쓴다


def _units(text):
    """한글은 라틴 글자의 두 배 폭으로 센다(줄 수 추정용)."""
    return sum(2 if ("가" <= c <= "힣" or "ㄱ" <= c <= "ㆎ") else 1 for c in text)


def _cell_lines(text, w_in, sz):
    """이 칸이 «몇 줄»이 되나 — 넣은 줄바꿈 + 폭에 따른 되감김."""
    cap = max(8, int(w_in / (sz * 0.0072)))         # 폭(인치)당 들어가는 반각 글자 수
    n = 0
    for seg in str(text).split("\n"):
        n += max(1, -(-_units(seg) // cap))          # ceil
    return n


def table(s, headers, rows, col_x, col_w, y0=1.72, sz=16, note=None):
    """가로줄만 있는 표. 반환값 = 닫는 줄(또는 note) 아래 y.

    ⚠ 행 높이는 «그 행의 실제 줄 수»로 잡는다. 고정 높이로 두면 여러 줄 셀이 다음 행을
    파고들어 글자가 겹친다 — 2026-09-11 실측이고, 좌표 검사로는 안 잡히고 렌더에서만 보인다.

    ⚠ **가로줄은 `col_x[0]` 부터 «항상 12.6 까지» 통짜로 그어진다** — 마지막 열이 어디서
    끝나든 상관없다. 그래서 표 오른쪽에 주석 상자를 두면 줄 위에 얹힌다(2026-09-12 실측).
    표 옆에 무언가를 놓고 싶으면 그 말을 표 «아래»로 내리거나 표를 쓰지 말 것.
    """
    L, W = col_x[0], 12.6 - col_x[0]
    line_h = sz * 0.0178 + 0.055                     # 한 줄이 차지하는 세로(인치)
    hdr_n = max(_cell_lines(h, col_w[j], sz) for j, h in enumerate(headers))
    hdr_h = max(0.42, hdr_n * line_h + 0.10)
    rect(s, L, y0, W, 0.035, INK)
    for j, hcell in enumerate(headers):
        tb(s, col_x[j], y0 + 0.10, col_w[j], hdr_h, hcell, sz=sz, bold=True, color=INK)
    rect(s, L, y0 + hdr_h + 0.20, W, 0.02, INK)
    y = y0 + hdr_h + 0.32
    for i, row in enumerate(rows):
        n = max(_cell_lines(c, col_w[j], sz) for j, c in enumerate(row))
        rh = max(0.40, n * line_h + 0.14)
        for j, cell in enumerate(row):
            tb(s, col_x[j], y, col_w[j], rh, str(cell), sz=sz, color=BODY)
        if i < len(rows) - 1:
            rect(s, L, y + rh - 0.03, W, 0.012, HAIRLINE)
        y += rh
    rect(s, L, y, W, 0.035, INK)
    if note:
        nh = max(0.34, _cell_lines(note, W, 14) * (14 * 0.0178 + 0.055) + 0.10)
        tb(s, L, y + 0.14, W, nh, note, sz=14, color=MUTE)
        y += 0.14 + nh
    return y


def code(s, l, t, w, h, lines, sz=15):
    """명령·경로 블록. 회색 바탕 + Consolas."""
    rect(s, l, t, w, h, CODEBG)
    bx = _flat(s.shapes.add_textbox(Inches(l + 0.18), Inches(t + 0.10),
                                    Inches(w - 0.30), Inches(h - 0.20)))
    tf = bx.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln if ln else " "
        p.font.size, p.font.name, p.font.color.rgb = Pt(sz), MONO, BODY
        p.line_spacing = 1.15
    return t + h


def picture(s, path, l, t, max_w, max_h, cap=None, border=True):
    """그림 한 장을 (max_w x max_h) 안에 비율 그대로 넣는다. 반환 = (캡션까지 포함한) 아래 y.

    ⚠ 화면 캡처는 «테두리»가 있어야 슬라이드 바탕과 구분된다(흰 배경이라 안 그러면 떠 보인다).
    """
    from PIL import Image
    iw, ih = Image.open(path).size
    ar = iw / ih
    w, h = max_w, max_w / ar
    if h > max_h:
        h, w = max_h, max_h * ar
    left = l + (max_w - w) / 2.0
    pic = _flat(s.shapes.add_picture(str(path), Inches(left), Inches(t), Inches(w), Inches(h)))
    if border:
        pic.line.color.rgb = HAIRLINE
        pic.line.width = Pt(0.75)
    y = t + h
    if cap:
        tb(s, l, y + 0.10, max_w, 0.40, cap, sz=14, color=MUTE, align=PP_ALIGN.CENTER)
        y += 0.50
    return y


def note(s, text, y=6.86, sz=14):
    return tb(s, 0.78, y, 11.9, 0.45, text, sz=sz, color=MUTE)


def new_deck_mono(template=None):
    return _new_deck(template)


# 짧은 이름으로도 쓴다
new_deck = new_deck_mono
